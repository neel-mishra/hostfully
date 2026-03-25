from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

from pacing_allocation import ChannelPacingSummary
from model_legend import MODEL_LEGEND_LINES


def write_pacing_workbook(
    path: Path,
    as_of: date,
    currency: str,
    meta_df: pd.DataFrame,
    meta_summary: ChannelPacingSummary,
    google_df: pd.DataFrame,
    google_summary: ChannelPacingSummary,
    total_target_daily: float,
    portfolio_warnings: Optional[List[str]] = None,
) -> Path:
    """
    One sheet: portfolio KPI block, Meta table, Google table.
    Requires openpyxl.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise ImportError("Install openpyxl: pip install openpyxl") from e

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget pacing"
    bold = Font(bold=True)

    row = 1
    ws.cell(row=row, column=1, value="Paid ads budget pacing").font = bold
    row += 1
    ws.cell(row=row, column=1, value="As-of date:")
    ws.cell(row=row, column=2, value=as_of.isoformat())
    row += 1
    ws.cell(row=row, column=1, value="Remaining days")
    ws.cell(row=row, column=2, value=max(meta_summary.days_remaining, google_summary.days_remaining))
    row += 1
    ws.cell(row=row, column=1, value="Currency:")
    ws.cell(row=row, column=2, value=currency)
    row += 1
    ws.cell(
        row=row,
        column=1,
        value="Sign: pacing delta = monthly budget − forecast month spend (positive = underspend).",
    )
    row += 1
    ws.cell(row=row, column=1, value="Legend").font = bold
    row += 1
    for line in MODEL_LEGEND_LINES:
        ws.cell(row=row, column=1, value=line)
        row += 1
    row += 1

    # Portfolio KPI (both channels)
    port = _portfolio_kpis(meta_summary, google_summary, total_target_daily)
    ws.cell(row=row, column=1, value="Portfolio (Meta + Google)").font = bold
    row += 1
    for label, val in port:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=val)
        row += 1

    if portfolio_warnings:
        row += 1
        ws.cell(row=row, column=1, value="Checks:").font = bold
        row += 1
        for w in portfolio_warnings:
            ws.cell(row=row, column=1, value=w)
            row += 1

    row += 1

    # Meta
    row = _write_table(ws, row, "Meta", meta_df, meta_summary, bold)
    row += 2

    # Google
    _write_table(ws, row, "Google", google_df, google_summary, bold)

    # Column widths (rough)
    for col in range(1, 19):
        ws.column_dimensions[get_column_letter(col)].width = 18

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _portfolio_kpis(
    meta: ChannelPacingSummary,
    google: ChannelPacingSummary,
    total_target_daily: float,
) -> List[Tuple[str, Any]]:
    B = meta.monthly_budget_total + google.monthly_budget_total
    S = meta.spend_mtd_total + google.spend_mtd_total
    F = meta.forecast_month_spend + google.forecast_month_spend
    cur_d = meta.sum_current_daily_budget + google.sum_current_daily_budget
    pacing_delta = B - F
    d_rem = max(meta.days_remaining, google.days_remaining)
    req = (B - S) / d_rem if d_rem > 0 else 0.0
    return [
        ("Total monthly budget", round(B, 2)),
        ("Total spent MTD (all campaigns in pull)", round(S, 2)),
        ("Total forecast month spend", round(F, 2)),
        ("Total pacing delta (monthly − forecast)", round(pacing_delta, 2)),
        ("Total current daily budget (all campaigns in pull)", round(cur_d, 2)),
        ("Total target daily budget (sum of recommended)", round(total_target_daily, 2)),
        ("Required daily run rate to hit plan (portfolio)", round(req, 2)),
    ]


def _write_table(
    ws,
    start_row: int,
    title: str,
    df: pd.DataFrame,
    summary: ChannelPacingSummary,
    bold,
) -> int:
    row = start_row
    ws.cell(row=row, column=1, value=f"{title} — campaigns").font = bold
    row += 1

    cols = [
        "source",
        "campaign_id",
        "platform_campaign_name",
        "campaign_name",
        "current_daily_budget",
        "monthly_budget",
        "spend_mtd",
        "pipeline_amount",
        "closed_won_amount",
        "sql_count",
        "forecast_month_spend",
        "pacing_delta_campaign",
        "recommended_daily_budget_capped",
        "daily_budget_delta",
        "pct_budget_change",
        "allocation_share_pct",
        "note",
    ]
    rename = {
        "source": "Source",
        "campaign_id": "Playbook campaign_id",
        "platform_campaign_name": "Platform campaign name",
        "campaign_name": "Playbook campaign_name",
        "current_daily_budget": "Current daily budget",
        "monthly_budget": "Monthly budget",
        "spend_mtd": "Spent MTD",
        "pipeline_amount": "Pipeline $",
        "closed_won_amount": "Closed won $",
        "sql_count": "SQL deals",
        "forecast_month_spend": "Forecast month spend",
        "pacing_delta_campaign": "Overspend / underspend delta",
        "recommended_daily_budget_capped": "Recommended new daily budget",
        "daily_budget_delta": "Daily budget delta",
        "pct_budget_change": "% budget change (new vs current)",
        "allocation_share_pct": "% of total daily budget delta",
        "note": "Note",
    }

    col_index = {name: i + 1 for i, name in enumerate(cols)}

    for c, col in enumerate(cols, start=1):
        ws.cell(row=row, column=c, value=rename.get(col, col)).font = bold
    row += 1

    if df is not None and not df.empty:
        display = df[[c for c in cols if c in df.columns]].copy()
        for _, r in display.iterrows():
            for c, col in enumerate(cols, start=1):
                if col not in display.columns:
                    continue
                val = r[col]
                ws.cell(row=row, column=c, value=_excel_friendly(val))
            row += 1

    # Channel total row (align to numeric columns by name)
    ws.cell(row=row, column=1, value=f"{title} — TOTAL").font = bold
    if summary.monthly_budget_total or True:
        ws.cell(row=row, column=col_index["current_daily_budget"], value=round(summary.sum_current_daily_budget, 2))
        ws.cell(row=row, column=col_index["monthly_budget"], value=round(summary.monthly_budget_total, 2))
        ws.cell(row=row, column=col_index["spend_mtd"], value=round(summary.spend_mtd_total, 2))
        if "pipeline_amount" in col_index and df is not None and not df.empty:
            ws.cell(row=row, column=col_index["pipeline_amount"], value=round(float(df.get("pipeline_amount", pd.Series(dtype=float)).sum()), 2))
        if "closed_won_amount" in col_index and df is not None and not df.empty:
            ws.cell(row=row, column=col_index["closed_won_amount"], value=round(float(df.get("closed_won_amount", pd.Series(dtype=float)).sum()), 2))
        if "sql_count" in col_index and df is not None and not df.empty:
            ws.cell(row=row, column=col_index["sql_count"], value=round(float(df.get("sql_count", pd.Series(dtype=float)).sum()), 2))
        ws.cell(row=row, column=col_index["forecast_month_spend"], value=round(summary.forecast_month_spend, 2))
        pdelta = (
            float(df["pacing_delta_campaign"].sum())
            if df is not None and not df.empty and "pacing_delta_campaign" in df.columns
            else summary.pacing_delta
        )
        ws.cell(row=row, column=col_index["pacing_delta_campaign"], value=round(pdelta, 2))
        rec = df["recommended_daily_budget_capped"].sum() if df is not None and not df.empty else 0
        ws.cell(row=row, column=col_index["recommended_daily_budget_capped"], value=round(float(rec), 2))
        # Required channel daily delta (pacing delta / remaining days)
        ws.cell(row=row, column=col_index["daily_budget_delta"], value=round(float(summary.daily_pool), 2))
        if "pct_budget_change" in col_index and df is not None and not df.empty:
            cur = float(df["current_daily_budget"].sum()) if "current_daily_budget" in df.columns else 0.0
            rec_total = float(df["recommended_daily_budget_capped"].sum()) if "recommended_daily_budget_capped" in df.columns else 0.0
            pct = (100.0 * (rec_total - cur) / cur) if cur > 0 else 0.0
            ws.cell(row=row, column=col_index["pct_budget_change"], value=round(pct, 2))
        if "allocation_share_pct" in col_index:
            ws.cell(row=row, column=col_index["allocation_share_pct"], value=100.0)
        ws.cell(row=row, column=col_index["note"], value="Channel aggregate")
    row += 1
    return row


def _excel_friendly(val: Any) -> Any:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return ""
    if isinstance(val, str) and val.lower() == "nan":
        return ""
    return val


